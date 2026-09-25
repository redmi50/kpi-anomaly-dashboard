"""The dashboard artefact, and the model a BI tool binds to.

The CV for this project claims a dashboard that presents KPI reporting alongside
documented root causes and corrected figures. These tests check that the workbook
is produced, that it carries those three things, that every number in it came from
the warehouse, and that the model export is one file per table and per certified
view. A dashboard that restates a figure is a second place for it to be wrong, so
the workbook is read back and compared against the investigation rather than
against a written expectation.
"""

from __future__ import annotations

import csv

import pytest
from openpyxl import load_workbook

from kpi_audit.__main__ import main
from kpi_audit.catalogue import TABLE_NAMES, VIEW_NAMES
from kpi_audit.dashboard import build_workbook, export_model

SHEET_ORDER = (
    "KPI ledger",
    "Findings",
    "Certified monthly",
    "Certified channel",
    "Load reconciliation",
    "Access",
    "Warehouse",
)


@pytest.fixture
def workbook(report, connection, tmp_path):
    """The workbook, built once per test from the shared investigation."""
    path = build_workbook(report, connection, tmp_path / "kpi_dashboard.xlsx")
    return load_workbook(path)


def rows_of(sheet) -> list[list]:
    """The data rows of a sheet, with the title, note and headings removed."""
    return [
        [cell.value for cell in row]
        for row in sheet.iter_rows(min_row=4)
    ]


class TestWorkbook:
    def test_the_workbook_is_written_where_it_was_asked_for(self, report, connection, tmp_path):
        path = build_workbook(report, connection, tmp_path / "nested" / "kpi.xlsx")
        assert path.is_file()
        assert path.name == "kpi.xlsx"

    def test_the_sheets_are_in_reading_order(self, workbook):
        assert workbook.sheetnames == list(SHEET_ORDER)

    def test_every_sheet_states_what_it_holds(self, workbook):
        for name in SHEET_ORDER:
            sheet = workbook[name]
            assert sheet["A1"].value == name
            assert sheet["A2"].value, name

    def test_every_sheet_has_a_frozen_heading_row(self, workbook):
        for name in SHEET_ORDER:
            assert workbook[name].freeze_panes == "A4"


class TestLedgerSheet:
    def test_every_scenario_appears_once(self, report, workbook):
        ledger = rows_of(workbook["KPI ledger"])
        assert [row[0] for row in ledger] == [inv.scenario_id for inv in report.investigations]

    def test_the_reported_and_corrected_figures_are_the_measured_ones(self, report, workbook):
        ledger = {
            row[0]: row for row in rows_of(workbook["KPI ledger"])
        }
        for investigation in report.investigations:
            row = ledger[investigation.scenario_id]
            assert row[4] == investigation.measurement.reported
            assert row[5] == investigation.measurement.agreed

    def test_the_row_counts_are_the_measured_row_counts(self, report, workbook):
        ledger = {row[0]: row for row in rows_of(workbook["KPI ledger"])}
        investigation = report.by_id()["window_first_month_dropped"]
        row = ledger["window_first_month_dropped"]
        assert row[10] == len(investigation.measurement.reported_rows)
        assert row[11] == len(investigation.measurement.agreed_rows)

    def test_the_control_is_reported_as_having_no_severity(self, workbook):
        ledger = {row[0]: row for row in rows_of(workbook["KPI ledger"])}
        assert ledger["control_certified"][8] in ("", None)

    def test_the_discrepancy_is_stored_as_a_fraction(self, report, workbook):
        ledger = {row[0]: row for row in rows_of(workbook["KPI ledger"])}
        row = ledger["window_first_month_dropped"]
        assert row[6] == pytest.approx(
            report.by_id()["window_first_month_dropped"].measurement.relative
        )


class TestFindingsSheet:
    def test_the_findings_are_in_severity_order(self, report, workbook):
        findings = rows_of(workbook["Findings"])
        assert [row[1] for row in findings] == [
            finding.scenario_id for finding in report.sorted_findings
        ]

    def test_the_control_raises_no_row_on_the_findings_sheet(self, report, workbook):
        findings = rows_of(workbook["Findings"])
        assert "control_certified" not in [row[1] for row in findings]
        assert len(findings) == len(report.sorted_findings)

    def test_a_finding_carries_its_cause_and_correction(self, report, workbook):
        findings = {row[1]: row for row in rows_of(workbook["Findings"])}
        finding = report.by_id()["filter_returns_counted"].finding
        row = findings["filter_returns_counted"]
        assert row[8] == finding.cause
        assert row[9] == finding.correction

    def test_a_finding_carries_the_query_that_produced_the_number(self, report, workbook):
        findings = {row[1]: row for row in rows_of(workbook["Findings"])}
        row = findings["window_final_day_lost"]
        assert "SELECT" in row[10]
        assert "main.fact_order_lines" in row[10]


class TestCertifiedSheets:
    def test_the_monthly_sheet_has_a_row_per_month(self, connection, workbook):
        expected = connection.execute(
            "SELECT COUNT(*) FROM v_revenue_by_month"
        ).fetchone()[0]
        assert len(rows_of(workbook["Certified monthly"])) == expected

    def test_the_channel_sheet_reports_the_unmapped_member(self, workbook):
        sheet = workbook["Certified channel"]
        headings = [cell.value for cell in sheet[3]]
        channel_column = headings.index("channel_name")
        names = [row[channel_column] for row in rows_of(sheet)]
        assert "UNMAPPED" in names

    def test_the_reconciliation_sheet_shows_the_batch_that_disagrees(self, workbook):
        sheet = workbook["Load reconciliation"]
        headings = [cell.value for cell in sheet[3]]
        variance_column = headings.index("variance")
        variances = [row[variance_column] for row in rows_of(sheet)]
        assert sum(1 for variance in variances if variance) == 1

    def test_the_certified_total_matches_the_ledger(self, report, workbook):
        sheet = workbook["Certified monthly"]
        headings = [cell.value for cell in sheet[3]]
        revenue_column = headings.index("net_revenue")
        total = round(sum(row[revenue_column] for row in rows_of(sheet)), 2)
        assert total == pytest.approx(report.by_id()["control_certified"].measurement.agreed)


class TestAccessSheet:
    def test_every_role_and_column_appears(self, connection, workbook):
        sheet = workbook["Access"]
        probed = {
            (row[0], row[1], row[2]) for row in rows_of(sheet)
        }
        assert len(probed) > 0
        roles = {role for role, _, _ in probed}
        assert roles == {"engineer", "analyst", "auditor", "contractor"}

    def test_the_analyst_is_recorded_as_refused_the_list_price(self, workbook):
        sheet = workbook["Access"]
        cells = {(row[0], row[1], row[2]): row[3] for row in rows_of(sheet)}
        assert cells[("analyst", "dim_product", "unit_price")] is False
        assert cells[("engineer", "dim_product", "unit_price")] is True

    def test_the_matrix_is_not_all_allow_or_all_deny(self, workbook):
        decisions = {row[3] for row in rows_of(workbook["Access"])}
        assert decisions == {True, False}


class TestWarehouseSheet:
    def test_every_table_and_view_appears(self, workbook):
        objects = [row[0] for row in rows_of(workbook["Warehouse"])]
        assert objects == list(TABLE_NAMES) + list(VIEW_NAMES)

    def test_the_row_counts_are_read_from_the_warehouse(self, connection, workbook):
        counts = {row[0]: row[3] for row in rows_of(workbook["Warehouse"])}
        for name, expected in counts.items():
            actual = connection.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0]
            assert actual == expected, name


class TestModelExport:
    def test_one_file_per_table_and_per_view(self, connection, tmp_path):
        written = export_model(connection, tmp_path / "model")
        assert [path.name for path in written] == [
            f"{name}.csv" for name in list(TABLE_NAMES) + list(VIEW_NAMES)
        ]

    def test_every_file_exists(self, connection, tmp_path):
        written = export_model(connection, tmp_path / "model")
        assert all(path.is_file() for path in written)

    def test_the_fact_csv_has_one_row_per_loaded_line(self, connection, tmp_path):
        export_model(connection, tmp_path)
        with (tmp_path / "fact_order_lines.csv").open(encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        expected = connection.execute(
            "SELECT COUNT(*) FROM fact_order_lines"
        ).fetchone()[0]
        assert len(rows) == expected

    def test_the_fact_csv_names_every_catalogue_column(self, connection, tmp_path):
        from kpi_audit.catalogue import TABLES

        export_model(connection, tmp_path)
        with (tmp_path / "fact_order_lines.csv").open(encoding="utf-8") as handle:
            headings = next(csv.reader(handle))
        assert headings == list(TABLES["fact_order_lines"].column_names)

    def test_the_view_csv_holds_the_corrected_figures(self, connection, tmp_path):
        """A model binds to the corrected definition, not to a copied number."""
        from kpi_audit.warehouse import scalar

        export_model(connection, tmp_path)
        with (tmp_path / "v_revenue_by_month.csv").open(encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        exported = round(sum(float(row["net_revenue"]) for row in rows), 2)
        assert exported == pytest.approx(
            scalar(connection, "SELECT SUM(net_revenue) FROM v_revenue_by_month")
        )

    def test_the_load_log_is_exported_with_its_row_counts(self, connection, tmp_path):
        export_model(connection, tmp_path)
        with (tmp_path / "load_audit.csv").open(encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        assert "row_count" in rows[0]


class TestDashboardCommand:
    def test_dashboard_writes_the_workbook_and_the_model(self, tmp_path):
        assert main(["dashboard", "--out", str(tmp_path)]) == 0
        assert (tmp_path / "kpi_dashboard.xlsx").is_file()
        assert (tmp_path / "model").is_dir()

    def test_dashboard_prints_what_it_wrote(self, tmp_path, capsys):
        main(["dashboard", "--out", str(tmp_path)])
        out = capsys.readouterr().out
        assert "kpi_dashboard.xlsx" in out
        assert f"{len(TABLE_NAMES) + len(VIEW_NAMES)} model files" in out

    def test_dashboard_writes_the_sheets_it_printed(self, tmp_path):
        main(["dashboard", "--out", str(tmp_path)])
        workbook = load_workbook(tmp_path / "kpi_dashboard.xlsx")
        assert workbook.sheetnames == list(SHEET_ORDER)

    def test_dashboard_accepts_a_subset(self, tmp_path):
        assert main(
            ["dashboard", "--scenarios", "control_certified", "--out", str(tmp_path)]
        ) == 0
        workbook = load_workbook(tmp_path / "kpi_dashboard.xlsx")
        assert len(rows_of(workbook["KPI ledger"])) == 1
        assert len(rows_of(workbook["Findings"])) == 0

    def test_dashboard_rejects_an_unknown_scenario(self, tmp_path, capsys):
        assert main(["dashboard", "--scenarios", "nope", "--out", str(tmp_path)]) == 2
        assert "Unknown scenarios" in capsys.readouterr().err
        assert not (tmp_path / "kpi_dashboard.xlsx").exists()
