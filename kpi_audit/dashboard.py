"""The dashboard artefact.

The CV for this project claims a dashboard that presents KPI reporting alongside
documented root causes and corrected figures. This module produces that artefact
rather than describing it.

Two outputs, for two purposes.

* An Excel workbook, because that is a format a person can open and read. It
  carries the KPI ledger, the findings with their root causes, the certified
  figures, the load reconciliation and the grants matrix on separate sheets.
* A star schema exported as CSV, one file per table, plus the certified views.
  Power BI binds to that folder as a data model; the measures it needs are the
  figures in the workbook, which are already computed from the certified views,
  so the dashboard and the audit cannot disagree about what the number is.

Everything written here is read from the warehouse or from the investigation.
Nothing is typed in, which is the point: a dashboard that restates a number is a
second place for that number to be wrong.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .access import grants_report
from .anomalies import ANOMALY_IDS
from .catalogue import TABLES, VIEW_NAMES
from .investigation import InvestigationReport
from .warehouse import rows_as_dicts

# The workbook is built sheet by sheet in the order a reader works through it:
# what was reported, what was wrong with it, what the corrected definitions say,
# and what the load log says about the batches behind it.
BOLD = Font(bold=True)


def build_workbook(
    report: InvestigationReport,
    connection,
    path: str | Path,
) -> Path:
    """Write the dashboard workbook and return its path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    # The default sheet is removed rather than used, so that the sheet order is
    # the one stated above rather than an accident of construction.
    workbook.remove(workbook.active)

    _ledger_sheet(workbook, report)
    _findings_sheet(workbook, report)
    _table_sheet(workbook, "Certified monthly", _view("v_revenue_by_month", connection))
    _table_sheet(workbook, "Certified channel", _view("v_revenue_by_channel", connection))
    _table_sheet(workbook, "Load reconciliation", _view("v_load_reconciliation", connection))
    _access_sheet(workbook, connection)
    _warehouse_sheet(workbook, connection)

    workbook.save(path)
    return path


def export_model(connection, out_dir: str | Path) -> list[Path]:
    """Export the star schema as CSV, for a Power BI model to bind to.

    One file per table and one per certified view. The views are exported
    alongside the tables because a Power BI model should bind to the corrected
    definition rather than to a measure written into the report, which is how a
    correction ends up applied in one place and not another.
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    for name in list(TABLES) + list(VIEW_NAMES):
        rows = rows_as_dicts(connection, f"SELECT * FROM {name}")
        path = out_dir / f"{name}.csv"
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            if rows:
                writer.writerow(list(rows[0]))
                for row in rows:
                    writer.writerow([row[column] for column in rows[0]])
            else:
                columns = (
                    TABLES[name].column_names if name in TABLES else ()
                )
                writer.writerow(columns)
        written.append(path)

    return written


def _view(name: str, connection) -> list[dict[str, object]]:
    """Read a certified view."""
    return rows_as_dicts(connection, f"SELECT * FROM {name}")


def _ledger_sheet(workbook: Workbook, report: InvestigationReport) -> None:
    """The reported figure, the corrected figure, and the size of the gap."""
    headings = (
        "Scenario",
        "Family",
        "Question",
        "Measure",
        "Reported",
        "Corrected",
        "Discrepancy",
        "Direction",
        "Severity",
        "Changed clause",
        "Rows reported",
        "Rows corrected",
        "Wrongly included",
        "Wrongly excluded",
    )
    rows: list[list[object]] = []
    for investigation in report.investigations:
        measurement = investigation.measurement
        scenario = investigation.scenario
        rows.append(
            [
                scenario.scenario_id,
                scenario.family,
                scenario.question,
                scenario.definition.measure,
                measurement.reported,
                measurement.agreed,
                measurement.relative,
                measurement.direction,
                "" if investigation.finding is None else investigation.finding.severity,
                scenario.changed_clause,
                len(measurement.reported_rows),
                len(measurement.agreed_rows),
                len(measurement.wrongly_included),
                len(measurement.wrongly_excluded),
            ]
        )
    note = (
        "Discrepancy is the distance from the corrected figure, as a fraction of it. "
        "The row counts are the rows each definition aggregated."
    )
    _write_sheet(workbook, "KPI ledger", headings, rows, note, percent_columns=(7,))
    sheet = workbook["KPI ledger"]
    # The ledger is the sheet a reader scans, so the two value columns are made
    # readable as values rather than left to the reader to widen.
    for row in sheet.iter_rows(min_row=4, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "#,##0.00"


def _findings_sheet(workbook: Workbook, report: InvestigationReport) -> None:
    """One row per finding, with the cause and the correction."""
    headings = (
        "Severity",
        "Scenario",
        "Title",
        "Family",
        "Changed clause",
        "Reported",
        "Corrected",
        "Discrepancy",
        "Root cause",
        "Correction",
        "Reporting query",
    )
    rows: list[list[object]] = []
    for finding in report.sorted_findings:
        evidence = finding.evidence
        rows.append(
            [
                finding.severity,
                finding.scenario_id,
                finding.title,
                finding.family,
                evidence.get("changed_clause", ""),
                evidence.get("reported"),
                evidence.get("agreed"),
                evidence.get("relative"),
                finding.cause,
                finding.correction,
                str(evidence.get("reporting_sql", "")).strip(),
            ]
        )
    note = (
        "Findings are ordered by severity then by scenario id. The cause and the "
        "correction are the ones the investigation produced; the two value columns "
        "are the measurements that produced the severity."
    )
    _write_sheet(workbook, "Findings", headings, rows, note, percent_columns=(8,))


def _access_sheet(workbook: Workbook, connection) -> None:
    """The grants matrix, as probed rather than as documented."""
    headings = ("Role", "Table", "Column", "Read allowed")
    rows: list[list[object]] = []
    for entry in grants_report(connection):
        for cell in entry["cells"]:
            rows.append(
                [entry["role"], cell["table"], cell["column"], bool(cell["allowed"])]
            )
    note = (
        "Each row is one column and the answer the engine gave when the role was "
        "asked for it. This matrix is produced by attempting every read rather than "
        "by restating the policy, so the two cannot drift."
    )
    _write_sheet(workbook, "Access", headings, rows, note)


def _warehouse_sheet(workbook: Workbook, connection) -> None:
    """What is in the warehouse, and how many defects the audit is looking for."""
    rows: list[list[object]] = []
    for name, table in TABLES.items():
        count = connection.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0]
        rows.append([name, table.kind, table.grain, count])
    for name in VIEW_NAMES:
        count = connection.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0]
        rows.append([name, "certified view", "the corrected definition", count])
    note = (
        f"Generated from a fixed seed. {len(ANOMALY_IDS)} scenarios are investigated "
        "against this warehouse."
    )
    _write_sheet(
        workbook,
        "Warehouse",
        ("Object", "Kind", "Grain", "Rows"),
        rows,
        note,
    )


def _table_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, object]],
) -> None:
    """Write one sheet from a list of row dictionaries."""
    if not rows:
        _write_sheet(workbook, title, ("empty",), [], "The query returned no rows.")
        return
    headings = tuple(rows[0])
    _write_sheet(
        workbook,
        title,
        headings,
        [[row[heading] for heading in headings] for row in rows],
        "Read from the certified view rather than restated in the workbook.",
    )
    sheet = workbook[title]
    for position, heading in enumerate(headings, start=1):
        if heading in {"net_revenue", "unit_price", "gross_amount", "discount_amount"}:
            for row in sheet.iter_rows(min_row=4, min_col=position, max_col=position):
                for cell in row:
                    cell.number_format = "#,##0.00"


def _write_sheet(
    workbook: Workbook,
    title: str,
    headings: tuple,
    rows: list[list[object]],
    note: str,
    *,
    percent_columns: tuple[int, ...] = (),
) -> None:
    """Write one sheet: a title, a note, the headings, then the rows."""
    sheet = workbook.create_sheet(title)
    sheet.append([title])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append([note])
    sheet.append(list(headings))
    for cell in sheet[3]:
        cell.font = BOLD

    for row in rows:
        sheet.append(list(row))

    for index, heading in enumerate(headings, start=1):
        width = max(
            [len(str(heading))]
            + [
                len(_text(row[index - 1]))
                for row in rows[:200]
            ]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(10, width + 2))
        if index in percent_columns:
            for excel_row in sheet.iter_rows(
                min_row=4, min_col=index, max_col=index, max_row=sheet.max_row
            ):
                for cell in excel_row:
                    cell.number_format = "0.00%"
        if heading in {"Root cause", "Correction", "Reporting query", "note"}:
            for excel_row in sheet.iter_rows(
                min_row=4, min_col=index, max_col=index, max_row=sheet.max_row
            ):
                for cell in excel_row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A4"


def _text(value: object) -> str:
    """One value as text, for sizing a column."""
    if value is None:
        return ""
    return str(value)
